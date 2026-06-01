from __future__ import annotations

import csv
import io
import json
import re
import subprocess
import sys
from pathlib import Path
from typing import Any

from backend.app.services.account_names import account_name_for
from backend.app.services.account_names import load_account_name_map
from backend.app.services.artifacts import find_latest_artifact
from backend.app.services.artifacts import read_json
from roibang_v2.config import load_json
from roibang_v2.db.bootstrap import bootstrap_database
from roibang_v2.runs import write_run_artifact
from roibang_v2.workflows.ai_create_template_drafts import run_ai_create_template_drafts_request
from roibang_v2.workflows.product_automation_job import SUPPORTED_JOBS
from roibang_v2.workflows.product_automation_job import load_product_configs
from roibang_v2.workflows.product_automation_job import run_product_automation_job
from roibang_v2.workflows.product_automation_job import workflow_for_job

PRODUCT_KEY_PATTERN = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_-]{1,80}$")
ALLOWED_ACCOUNT_TEMPLATE_HEADERS = ["产品名", "产品 Key", "账户名", "账户 ID", "是否启用", "备注"]
ALLOWED_ACCOUNT_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "product": ("产品名", "产品", "product", "product_name"),
    "product_key": ("产品 Key", "产品Key", "product_key", "productKey"),
    "account_name": ("账户名", "账户名称", "account_name", "advertiser_name", "name"),
    "advertiser_id": ("账户 ID", "账户ID", "账户id", "advertiser_id", "account_id", "adv_id"),
    "enable": ("是否启用", "启用", "enable", "enabled", "status"),
    "account_remark": ("备注", "账户备注", "account_remark", "remark", "notes"),
}

JOB_INFO: dict[str, dict[str, str]] = {
    "material_daily_sync": {
        "label": "每日素材明细同步",
        "description": "按产品账户发现关键词找到有消耗账户，拉取素材日维度数据并写入本地数据库。",
    },
    "operation_log_sync": {
        "label": "操作日志同步",
        "description": "按有消耗账户同步平台操作记录，用于后续判断动作前后的效果。",
    },
    "daily_report_sync": {
        "label": "每日报表同步",
        "description": "同步产品级日报、项目和单元数据，沉淀本地学习与报表快照。",
    },
    "source_material_auto_push": {
        "label": "源素材账户自动补材",
        "description": "把产品有消耗账户里的新视频素材补进该产品源素材账户。",
    },
    "source_material_preload": {
        "label": "源素材预推送",
        "description": "把源素材账户的视频预推送到允许创建账户名单里的账户；新产品默认沿用点点英雄口径。",
    },
    "source_material_rollup": {
        "label": "源素材表现汇总重建",
        "description": "基于本地素材表现表重建源素材多窗口表现汇总，供创建和分析使用。",
    },
    "delivery_patrol": {
        "label": "小时投放巡检",
        "description": "按产品账户备注做小时级只读巡检，作为运营补充信息。",
    },
}


def build_product_automation_overview(
    *,
    project_root: str | Path,
    configs_dir: str | Path,
    runs_dir: str | Path,
) -> dict[str, Any]:
    configs_path = Path(configs_dir)
    products = load_product_configs(configs_path / "products")
    account_names = load_account_name_map(configs_path)
    latest_by_job = {job: _latest_product_job_results(runs_dir, job) for job in JOB_INFO}
    rows: list[dict[str, Any]] = []
    raw_products: list[dict[str, Any]] = []
    warnings: list[str] = []

    for product in products:
        product_key = _text(product.get("product_key"))
        product_name = _text(product.get("product"))
        enabled_jobs = _enabled_jobs(product)
        source_advertiser_id = _text(product.get("source_advertiser_id"))
        source_advertiser_name = _source_account_name(product, account_names)
        allowed_path = _text(product.get("allowed_target_accounts_path"))
        allowed_count, allowed_warning = _allowed_account_count(project_root, allowed_path)
        if allowed_warning:
            warnings.append(f"{product_name or product_key}：{allowed_warning}")
        preload_scope = _preload_target_scope(product)
        latest_status = _latest_status_text(product_key, latest_by_job)
        rows.append(
            {
                "产品": product_name,
                "产品 Key": product_key,
                "源素材账户名": source_advertiser_name,
                "源素材账户 ID": source_advertiser_id,
                "允许创建账户": allowed_count,
                "账户发现关键词": _account_keyword(product),
                "预推送目标": _preload_scope_label(preload_scope),
                "启用定时任务": "、".join(JOB_INFO[job]["label"] for job in enabled_jobs) if enabled_jobs else "未启用",
                "最近运行": latest_status,
            }
        )
        raw_products.append(_product_form_payload(product, source_advertiser_name, enabled_jobs))

    return {
        "summary": {
            "title": "产品自动化配置",
            "status": "loaded",
            "risk_level": "medium" if warnings else "low",
            "execution_enabled": False,
            "items": [
                {"label": "产品数", "value": len(rows)},
                {"label": "任务类型", "value": len(JOB_INFO)},
                {"label": "预演能力", "value": "只生成 JSON 和命令，不执行真实动作"},
            ],
            "warnings": warnings,
            "blocking_reasons": [],
        },
        "table": {
            "columns": [
                "产品",
                "产品 Key",
                "源素材账户名",
                "源素材账户 ID",
                "允许创建账户",
                "账户发现关键词",
                "预推送目标",
                "启用定时任务",
                "最近运行",
            ],
            "rows": rows,
        },
        "sections": [
            {
                "title": "定时任务说明",
                "table": {
                    "columns": ["任务", "配置键", "中文说明"],
                    "rows": [
                        {"任务": info["label"], "配置键": job, "中文说明": info["description"]}
                        for job, info in JOB_INFO.items()
                    ],
                },
            }
        ],
        "artifact_path": "",
        "raw": {
            "products": raw_products,
            "jobs": _job_catalog(),
        },
    }


def save_product_automation_config(
    *,
    configs_dir: str | Path,
    body: dict[str, Any],
) -> dict[str, Any]:
    blocking = _validate_save_request(body)
    if blocking:
        return _blocked_result("产品配置未保存", blocking, body)

    product_key = _text(body.get("product_key"))
    products_dir = Path(configs_dir) / "products"
    path = products_dir / f"{product_key}.local.json"
    existing = read_json(path)
    enabled_jobs = [job for job in body.get("enabled_jobs", []) if job in JOB_INFO]
    automation = _updated_automation(existing, body, enabled_jobs)
    payload = {
        **existing,
        "product_key": product_key,
        "product": _text(body.get("product")),
        "platform": _text(body.get("platform") or existing.get("platform") or "WECHAT_GAME"),
        "source_advertiser_id": _text(body.get("source_advertiser_id")),
        "source_advertiser_name": _text(body.get("source_advertiser_name")),
        "organization_id": _text(body.get("organization_id")),
        "allowed_target_accounts_path": _text(body.get("allowed_target_accounts_path")),
        "automation": automation,
    }
    products_dir.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    rows = [
        {"配置项": "产品", "值": payload["product"]},
        {"配置项": "产品 Key", "值": payload["product_key"]},
        {"配置项": "源素材账户名", "值": payload["source_advertiser_name"]},
        {"配置项": "源素材账户 ID", "值": payload["source_advertiser_id"]},
        {"配置项": "允许创建账户名单", "值": payload["allowed_target_accounts_path"]},
        {"配置项": "账户发现关键词", "值": automation["account_discovery"]["account_name_keyword"]},
        {"配置项": "源素材预推送目标", "值": "允许创建账户名单"},
        {"配置项": "启用定时任务", "值": "、".join(JOB_INFO[job]["label"] for job in enabled_jobs)},
    ]
    return {
        "summary": {
            "title": "产品配置已保存",
            "status": "committed",
            "risk_level": "low",
            "execution_enabled": False,
            "items": [
                {"label": "产品", "value": payload["product"]},
                {"label": "启用任务", "value": len(enabled_jobs)},
                {"label": "真实执行", "value": "未执行"},
            ],
            "warnings": ["保存只写本地产品配置；真实同步、补材和预推送仍由固定定时脚本执行。"],
            "blocking_reasons": [],
        },
        "table": {"columns": ["配置项", "值"], "rows": rows},
        "artifact_path": str(path),
        "raw": {"product": payload, "path": str(path)},
    }


def build_allowed_accounts_template(
    *,
    product_key: str,
    product: str,
    platform: str,
) -> tuple[bytes, str]:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.styles import PatternFill
    except ImportError as exc:
        raise RuntimeError("当前环境缺少 openpyxl，无法生成 Excel 模板") from exc

    product_key_text = _text(product_key) or "new-product"
    product_text = _text(product) or "填写产品名"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "允许创建账户名单"
    sheet.append(ALLOWED_ACCOUNT_TEMPLATE_HEADERS)
    sheet.append([product_text, product_key_text, "示例账户名", "1234567890", "是", "例如：郭靖账户"])
    sheet.append([product_text, product_key_text, "", "", "是", ""])
    sheet.freeze_panes = "A2"

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    widths = [18, 22, 28, 24, 12, 28]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    filename_key = product_key_text if PRODUCT_KEY_PATTERN.fullmatch(product_key_text) else "new-product"
    return buffer.getvalue(), f"allowed-create-accounts-template.{filename_key}.xlsx"


def import_allowed_accounts_upload(
    *,
    configs_dir: str | Path,
    filename: str,
    content: bytes,
    product_key: str,
    product: str,
    platform: str,
) -> dict[str, Any]:
    context_product_key = _text(product_key)
    context_product = _text(product)
    context_platform = _text(platform) or "WECHAT_GAME"
    blocking = []
    if not context_product_key:
        blocking.append("产品 Key 不能为空，请先填写产品 Key 再上传账户名单")
    elif not PRODUCT_KEY_PATTERN.fullmatch(context_product_key):
        blocking.append("产品 Key 只能包含英文、数字、中划线或下划线，且不能包含路径符号")
    if not context_product:
        blocking.append("产品名不能为空，请先填写产品名再上传账户名单")
    if blocking:
        return _allowed_accounts_result(
            title="允许创建账户名单未生成",
            status="blocked",
            rows=[],
            counts={"read": 0, "enabled": 0, "skipped": 0, "error": len(blocking)},
            blocking_reasons=blocking,
            artifact_path="",
            raw={"product_key": context_product_key, "product": context_product},
        )

    try:
        uploaded_rows = _parse_allowed_accounts_upload(filename, content)
    except RuntimeError as exc:
        return _allowed_accounts_result(
            title="允许创建账户名单未生成",
            status="blocked",
            rows=[],
            counts={"read": 0, "enabled": 0, "skipped": 0, "error": 1},
            blocking_reasons=[str(exc)],
            artifact_path="",
            raw={"filename": filename},
        )

    table_rows: list[dict[str, str]] = []
    enabled_accounts: list[dict[str, Any]] = []
    blocking_reasons: list[str] = []
    seen_ids: set[str] = set()
    counts = {"read": len(uploaded_rows), "enabled": 0, "skipped": 0, "error": 0}

    for index, row in enumerate(uploaded_rows, start=2):
        row_product = _field(row, "product") or context_product
        row_product_key = _field(row, "product_key") or context_product_key
        account_name = _field(row, "account_name")
        advertiser_id = _field(row, "advertiser_id")
        account_remark = _field(row, "account_remark")
        enabled = _enabled_for_import(_field(row, "enable"))
        issues = []
        action = "写入"
        if not enabled:
            action = "跳过"
            counts["skipped"] += 1
        else:
            if row_product_key != context_product_key:
                issues.append(f"产品 Key 与当前页面不一致：{row_product_key}")
            if row_product != context_product:
                issues.append(f"产品名与当前页面不一致：{row_product}")
            if not account_name:
                issues.append("账户名不能为空")
            if not advertiser_id:
                issues.append("账户 ID 不能为空")
            if advertiser_id and advertiser_id in seen_ids:
                issues.append("同次上传中账户 ID 重复")

            if issues:
                action = "错误"
                counts["error"] += 1
                blocking_reasons.extend(f"第 {index} 行：{issue}" for issue in issues)
            else:
                seen_ids.add(advertiser_id)
                counts["enabled"] += 1
                enabled_accounts.append(
                    {
                        "advertiser_id": advertiser_id,
                        "account_name": account_name,
                        "account_remark": account_remark,
                        "product": context_product,
                        "product_key": context_product_key,
                        "platform": context_platform,
                        "enable": True,
                    }
                )

        table_rows.append(
            {
                "处理方式": action,
                "产品": row_product,
                "产品 Key": row_product_key,
                "账户名": account_name,
                "账户 ID": advertiser_id,
                "是否启用": "是" if enabled else "否",
                "备注": account_remark,
                "问题": "；".join(issues),
            }
        )

    if not uploaded_rows:
        counts["error"] += 1
        blocking_reasons.append("上传表格没有读取到账户行")
    if not enabled_accounts and not blocking_reasons:
        blocking_reasons.append("上传表格里没有启用的账户")
    if blocking_reasons:
        return _allowed_accounts_result(
            title="允许创建账户名单未生成",
            status="blocked",
            rows=table_rows,
            counts=counts,
            blocking_reasons=blocking_reasons,
            artifact_path="",
            raw={"filename": filename, "product_key": context_product_key, "product": context_product},
        )

    relative_path = f"configs/allowed-create-accounts.{context_product_key}.local.json"
    path = Path(configs_dir) / f"allowed-create-accounts.{context_product_key}.local.json"
    payload = {
        "product": context_product,
        "product_key": context_product_key,
        "platform": context_platform,
        "allowed_target_accounts": enabled_accounts,
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return _allowed_accounts_result(
        title="允许创建账户名单已生成",
        status="committed",
        rows=table_rows,
        counts=counts,
        blocking_reasons=[],
        artifact_path=str(path),
        raw={
            "allowed_target_accounts_path": relative_path,
            "path": str(path),
            "accounts": enabled_accounts,
            "product_key": context_product_key,
            "product": context_product,
        },
    )


def build_product_automation_dry_run(
    *,
    project_root: str | Path,
    configs_dir: str | Path,
    runs_dir: str | Path,
    product_key: str,
    job: str,
    target_date: str,
) -> dict[str, Any]:
    if job not in JOB_INFO:
        return _blocked_result("产品自动化预演不可用", [f"不支持的任务类型：{job}"], {"job": job})
    if not _text(product_key):
        return _blocked_result("产品自动化预演不可用", ["请选择产品"], {"job": job})

    result = run_product_automation_job(
        job=job,
        products_dir=Path(configs_dir) / "products",
        request_dir=Path(project_root) / "data" / "requests" / "product-automation",
        runs_dir=runs_dir,
        product_key=product_key,
        target_date=_text(target_date) or "yesterday",
        enable_readonly=True,
        execute=False,
        yes=False,
        dry_run=True,
    )
    rows = []
    for item in result.get("results", []):
        command = " ".join(str(part) for part in item.get("command", []))
        rows.append(
            {
                "产品": item.get("product"),
                "产品 Key": item.get("product_key"),
                "任务": JOB_INFO[job]["label"],
                "请求 JSON": item.get("request_path"),
                "脚本命令": command,
                "真实执行": "否",
            }
        )
    status = "planned" if rows else "empty"
    warnings = [] if rows else ["没有找到启用该任务的产品配置，请先保存产品配置并启用对应定时任务。"]
    return {
        "summary": {
            "title": "产品自动化预演",
            "status": status,
            "risk_level": "low",
            "execution_enabled": False,
            "items": [
                {"label": "任务", "value": JOB_INFO[job]["label"]},
                {"label": "产品数", "value": len(rows)},
                {"label": "目标日期", "value": _text(target_date) or "yesterday"},
                {"label": "真实执行", "value": "否"},
            ],
            "warnings": warnings,
            "blocking_reasons": [],
        },
        "table": {"columns": ["产品", "产品 Key", "任务", "请求 JSON", "脚本命令", "真实执行"], "rows": rows},
        "artifact_path": str(result.get("artifact_path") or ""),
        "raw": result,
    }


def build_ai_template_drafts(
    *,
    project_root: str | Path,
    configs_dir: str | Path,
    runs_dir: str | Path,
    product_key: str,
    max_drafts: int = 5,
) -> dict[str, Any]:
    product_key_text = _text(product_key)
    if not product_key_text:
        return _blocked_result("AI 模板草稿不可用", ["请选择产品"], {"product_key": product_key})

    products = load_product_configs(Path(configs_dir) / "products")
    product = next((item for item in products if _text(item.get("product_key")) == product_key_text), None)
    if not product:
        return _blocked_result("AI 模板草稿不可用", [f"未找到产品配置：{product_key_text}"], {"product_key": product_key_text})

    product_name = _text(product.get("product"))
    source_advertiser_id = _text(product.get("source_advertiser_id"))
    if not source_advertiser_id:
        return _blocked_result(
            "AI 模板草稿不可用",
            [f"{product_name or product_key_text} 未配置源素材账户 ID"],
            {"product_key": product_key_text},
        )

    root = Path(project_root)
    db_path = root / "data" / "roibang_v2.sqlite3"
    if not db_path.exists():
        return _blocked_result("AI 模板草稿不可用", [f"本地数据库不存在：{db_path}"], {"product_key": product_key_text})
    bootstrap_database(db_path)

    result = run_ai_create_template_drafts_request(
        {
            "ai_create_template_drafts": {
                "product": product_name,
                "source_advertiser_id": source_advertiser_id,
                "manual_mode_dir": str(Path(configs_dir) / "create-modes"),
                "max_drafts": max(int(max_drafts or 5), 0),
            }
        },
        db_path=db_path,
        runs_dir=runs_dir,
    )
    drafts = [draft for draft in result.get("drafts", []) if isinstance(draft, dict)]
    rows = [_ai_template_draft_row(draft, product_name=product_name) for draft in drafts]
    status = "draft_only" if drafts else ("blocked" if result.get("blocking_reasons") else "empty")
    return {
        "summary": {
            "title": "AI 模板草稿",
            "status": status,
            "risk_level": "medium" if drafts else "low",
            "execution_enabled": False,
            "items": [
                {"label": "产品", "value": product_name or product_key_text},
                {"label": "草稿数", "value": len(drafts)},
                {"label": "真实执行", "value": "否"},
            ],
            "warnings": [
                "AI 模板草稿只写 data/runs，不写人工固定模板；转正前不能被创建脚本直接使用。"
            ]
            if drafts
            else [],
            "blocking_reasons": [_text(reason) for reason in result.get("blocking_reasons", []) if _text(reason)],
        },
        "table": {
            "columns": [
                "草稿名",
                "产品",
                "草稿 Key",
                "基于人工模板",
                "候选素材",
                "消耗",
                "转化",
                "ROI",
                "差异",
                "风险",
                "状态",
            ],
            "rows": rows,
        },
        "artifact_path": _text(result.get("artifact_path")),
        "raw": result,
    }


def build_ai_template_draft_preview(
    *,
    project_root: str | Path,
    configs_dir: str | Path,
    runs_dir: str | Path,
    product_key: str,
    artifact_path: str,
    draft_key: str,
) -> dict[str, Any]:
    product_key_text = _text(product_key)
    artifact_path_text = _text(artifact_path)
    draft_key_text = _text(draft_key)
    blocking = []
    if not product_key_text:
        blocking.append("请选择产品")
    elif not PRODUCT_KEY_PATTERN.fullmatch(product_key_text):
        blocking.append("产品 Key 只能包含英文、数字、中划线或下划线，且不能包含路径符号")
    if not artifact_path_text:
        blocking.append("缺少 AI 模板草稿结果文件")
    if not draft_key_text:
        blocking.append("请选择要转正预览的草稿")
    elif not PRODUCT_KEY_PATTERN.fullmatch(draft_key_text):
        blocking.append("草稿 Key 只能包含英文、数字、中划线或下划线，且不能包含路径符号")
    if blocking:
        return _blocked_result(
            "AI 模板草稿转正预览不可用",
            blocking,
            {"product_key": product_key_text, "artifact_path": artifact_path_text, "draft_key": draft_key_text},
        )

    products = load_product_configs(Path(configs_dir) / "products")
    product = next((item for item in products if _text(item.get("product_key")) == product_key_text), None)
    if not product:
        return _blocked_result("AI 模板草稿转正预览不可用", [f"未找到产品配置：{product_key_text}"], {"product_key": product_key_text})

    source_path = _project_path(project_root, artifact_path_text)
    if not source_path.exists():
        return _blocked_result(
            "AI 模板草稿转正预览不可用",
            [f"AI 模板草稿结果文件不存在：{artifact_path_text}"],
            {"product_key": product_key_text, "artifact_path": artifact_path_text, "draft_key": draft_key_text},
        )
    source_payload = read_json(source_path)
    drafts = source_payload.get("drafts") if isinstance(source_payload.get("drafts"), list) else []
    draft = next(
        (item for item in drafts if isinstance(item, dict) and _text(item.get("draft_key")) == draft_key_text),
        None,
    )
    if not draft:
        return _blocked_result(
            "AI 模板草稿转正预览不可用",
            [f"草稿结果文件中没有找到草稿：{draft_key_text}"],
            {"product_key": product_key_text, "artifact_path": str(source_path), "draft_key": draft_key_text},
        )

    product_name = _text(product.get("product"))
    target_path = Path(configs_dir) / "create-modes" / product_key_text / f"{draft_key_text}.local.json"
    proposed_create_mode = dict(
        draft.get("proposed_create_mode") if isinstance(draft.get("proposed_create_mode"), dict) else {}
    )
    proposed_create_mode.update(
        {
            "mode_key": draft_key_text,
            "product_key": product_key_text,
            "product": product_name,
            "source_advertiser_id": _text(product.get("source_advertiser_id")),
            "organization_id": _text(product.get("organization_id")),
        }
    )
    preview = {
        "ok": True,
        "workflow": "ai_template_draft_promotion_preview",
        "phase": "preview",
        "status": "preview_only",
        "execution_enabled": False,
        "external_api_calls": 0,
        "product_key": product_key_text,
        "product": product_name,
        "draft_key": draft_key_text,
        "draft_name": _text(draft.get("draft_name")),
        "target_path": str(target_path),
        "target_path_exists": target_path.exists(),
        "source_artifact_path": str(source_path),
        "summary": {
            "title": "AI 模板草稿转正预览",
            "中文摘要": f"为产品 {product_name or product_key_text} 的草稿 {draft_key_text} 生成转正预览 JSON；不写入人工固定模板，不执行真实创建。",
            "产品": product_name or product_key_text,
            "草稿 Key": draft_key_text,
            "目标文件": str(target_path),
            "真实执行": "否",
        },
        "proposed_create_mode": proposed_create_mode,
        "evidence": draft.get("evidence") if isinstance(draft.get("evidence"), dict) else {},
        "differences_from_manual_template": _string_list(draft.get("differences_from_manual_template")),
        "risk_notes": _string_list(draft.get("risk_notes")),
        "manual_template_boundary": {
            "writes_manual_template": False,
            "generates_create_plan": False,
            "requires_human_promotion_script": True,
            "target_manual_template_path": str(target_path),
            "preview_output_only": "data/runs/ai_template_draft_preview",
        },
        "execution": {
            "enabled": False,
            "status": "preview_only",
            "real_business_action": False,
        },
        "actions": [],
    }
    artifact = write_run_artifact(runs_dir, "ai_template_draft_preview", preview)
    warnings = [
        "这里只生成转正预览 JSON，不写入 configs/create-modes（人工创建模式目录）。",
        "后续真正转正必须由单独固定脚本读取这个预览，并由人工确认后写入。",
    ]
    if target_path.exists():
        warnings.append("目标文件已经存在；本次预览不会覆盖。")
    row = {
        "草稿名": preview["draft_name"],
        "产品": product_name,
        "产品 Key": product_key_text,
        "草稿 Key": draft_key_text,
        "目标模式 Key": draft_key_text,
        "目标文件": str(target_path),
        "基于人工模板": _text(draft.get("base_manual_mode_key")),
        "差异": "；".join(preview["differences_from_manual_template"]),
        "风险": "；".join(preview["risk_notes"]),
        "目标文件已存在": "是" if target_path.exists() else "否",
        "状态": "preview_only",
    }
    return {
        "summary": {
            "title": "AI 模板草稿转正预览",
            "status": "planned",
            "risk_level": "medium",
            "execution_enabled": False,
            "items": [
                {"label": "产品", "value": product_name or product_key_text},
                {"label": "草稿", "value": draft_key_text},
                {"label": "目标文件已存在", "value": "是" if target_path.exists() else "否"},
                {"label": "真实执行", "value": "否"},
            ],
            "warnings": warnings,
            "blocking_reasons": [],
        },
        "table": {
            "columns": [
                "草稿名",
                "产品",
                "产品 Key",
                "草稿 Key",
                "目标模式 Key",
                "目标文件",
                "基于人工模板",
                "差异",
                "风险",
                "目标文件已存在",
                "状态",
            ],
            "rows": [row],
        },
        "artifact_path": str(artifact),
        "raw": {"preview": {**preview, "artifact_path": str(artifact)}},
    }


def build_ai_template_draft_promote(
    *,
    project_root: str | Path,
    configs_dir: str | Path,
    runs_dir: str | Path,
    product_key: str,
    preview_path: str,
    replace: bool = False,
) -> dict[str, Any]:
    product_key_text = _text(product_key)
    preview_path_text = _text(preview_path)
    blocking = []
    if not product_key_text:
        blocking.append("请选择产品")
    elif not PRODUCT_KEY_PATTERN.fullmatch(product_key_text):
        blocking.append("产品 Key 只能包含英文、数字、中划线或下划线，且不能包含路径符号")
    if not preview_path_text:
        blocking.append("缺少转正预览 JSON 文件")
    if blocking:
        return _blocked_result(
            "AI 模板草稿写入创建模式不可用",
            blocking,
            {"product_key": product_key_text, "preview_path": preview_path_text},
        )

    script_path = Path(__file__).resolve().parents[3] / "scripts" / "run_ai_template_draft_promote.py"
    command = [
        sys.executable,
        str(script_path),
        "--preview",
        preview_path_text,
        "--mode-dir",
        str(Path(configs_dir) / "create-modes"),
        "--runs-dir",
        str(Path(runs_dir)),
        "--product-key",
        product_key_text,
    ]
    if replace:
        command.append("--replace")
    completed = subprocess.run(
        command,
        cwd=Path(project_root),
        capture_output=True,
        text=True,
        check=False,
    )
    result = _parse_script_json(completed.stdout)
    if not result:
        return _blocked_result(
            "AI 模板草稿写入创建模式失败",
            ["固定脚本没有输出有效 JSON"],
            {
                "command": command,
                "return_code": completed.returncode,
                "stdout": completed.stdout,
                "stderr": completed.stderr,
            },
        )

    ok = bool(result.get("ok"))
    summary = result.get("summary") if isinstance(result.get("summary"), dict) else {}
    row = {
        "产品": _text(summary.get("product")),
        "产品 Key": _text(summary.get("product_key")),
        "草稿 Key": _text(summary.get("draft_key")),
        "目标文件": _text(summary.get("target_path") or result.get("target_path")),
        "固定脚本": "scripts/run_ai_template_draft_promote.py",
        "真实执行": "否",
        "状态": _text(result.get("status")),
    }
    return {
        "summary": {
            "title": "AI 模板草稿已写入创建模式" if ok else "AI 模板草稿写入被阻止",
            "status": "committed" if ok else "blocked",
            "risk_level": "low" if ok else "medium",
            "execution_enabled": False,
            "items": [
                {"label": "产品", "value": row["产品"] or row["产品 Key"]},
                {"label": "草稿", "value": row["草稿 Key"]},
                {"label": "目标文件", "value": row["目标文件"]},
                {"label": "真实执行", "value": "否"},
            ],
            "warnings": ["写入的是本地创建模式 JSON；没有生成创建计划，也没有执行真实投放。"] if ok else [],
            "blocking_reasons": [_text(reason) for reason in result.get("blocking_reasons", []) if _text(reason)],
        },
        "table": {
            "columns": ["产品", "产品 Key", "草稿 Key", "目标文件", "固定脚本", "真实执行", "状态"],
            "rows": [row],
        },
        "artifact_path": _text(result.get("artifact_path")),
        "raw": {
            "result": result,
            "command": command,
            "return_code": completed.returncode,
            "stdout": completed.stdout,
            "stderr": completed.stderr,
        },
    }


def _validate_save_request(body: dict[str, Any]) -> list[str]:
    reasons = []
    product_key = _text(body.get("product_key"))
    if not product_key:
        reasons.append("产品 Key 不能为空")
    elif not PRODUCT_KEY_PATTERN.fullmatch(product_key):
        reasons.append("产品 Key 只能包含英文、数字、中划线或下划线，且不能包含路径符号")
    for key, label in [
        ("product", "产品名"),
        ("source_advertiser_name", "源素材账户名"),
        ("source_advertiser_id", "源素材账户 ID"),
        ("organization_id", "组织 ID"),
        ("allowed_target_accounts_path", "允许创建账户名单文件路径"),
        ("account_name_keyword", "账户发现关键词"),
    ]:
        if not _text(body.get(key)):
            reasons.append(f"{label}不能为空")
    enabled_jobs = body.get("enabled_jobs")
    if not isinstance(enabled_jobs, list) or not [job for job in enabled_jobs if job in JOB_INFO]:
        reasons.append("至少启用一个定时任务")
    return reasons


def _updated_automation(existing: dict[str, Any], body: dict[str, Any], enabled_jobs: list[str]) -> dict[str, Any]:
    automation = dict(existing.get("automation") if isinstance(existing.get("automation"), dict) else {})
    automation["enabled"] = True
    discovery = dict(automation.get("account_discovery") if isinstance(automation.get("account_discovery"), dict) else {})
    discovery["account_name_keyword"] = _text(body.get("account_name_keyword"))
    discovery["account_remark_equals"] = _text(body.get("account_remark_equals"))
    automation["account_discovery"] = discovery
    for job in JOB_INFO:
        job_cfg = dict(automation.get(job) if isinstance(automation.get(job), dict) else {})
        job_cfg["enabled"] = job in enabled_jobs
        if job == "material_daily_sync":
            job_cfg.setdefault("source", "account_name_keyword")
            job_cfg.setdefault("min_spend", 0)
        if job == "source_material_auto_push":
            job_cfg.setdefault("from_spent_accounts", True)
        if job == "source_material_preload":
            job_cfg["target_scope"] = "allowed_accounts"
        if job == "source_material_rollup":
            job_cfg.setdefault("date_range", {"start": "2026-02-10", "end": "yesterday"})
            job_cfg.setdefault("windows", [1, 3, 7, 15, 30, "all"])
        if job == "delivery_patrol":
            job_cfg.setdefault("account_scope", "account_remark")
        automation[job] = job_cfg
    return automation


def _product_form_payload(product: dict[str, Any], source_advertiser_name: str, enabled_jobs: list[str]) -> dict[str, Any]:
    return {
        "product_key": _text(product.get("product_key")),
        "product": _text(product.get("product")),
        "platform": _text(product.get("platform") or "WECHAT_GAME"),
        "source_advertiser_name": source_advertiser_name,
        "source_advertiser_id": _text(product.get("source_advertiser_id")),
        "organization_id": _text(product.get("organization_id")),
        "allowed_target_accounts_path": _text(product.get("allowed_target_accounts_path")),
        "account_name_keyword": _account_keyword(product),
        "account_remark_equals": _account_remark(product),
        "enabled_jobs": enabled_jobs,
        "config_path": _text(product.get("_config_path")),
    }


def _latest_product_job_results(runs_dir: str | Path, job: str) -> dict[str, dict[str, Any]]:
    path = find_latest_artifact(runs_dir, workflow_for_job(job))
    payload = read_json(path) if path else {}
    results = payload.get("results") if isinstance(payload.get("results"), list) else []
    rows = {}
    for item in results:
        if isinstance(item, dict):
            product_key = _text(item.get("product_key"))
            if product_key:
                rows[product_key] = item
    return rows


def _latest_status_text(product_key: str, latest_by_job: dict[str, dict[str, dict[str, Any]]]) -> str:
    parts = []
    for job, rows in latest_by_job.items():
        item = rows.get(product_key)
        if not item:
            continue
        ok = bool(item.get("ok", True))
        parsed = item.get("parsed_stdout") if isinstance(item.get("parsed_stdout"), dict) else {}
        status = _text(parsed.get("status") or ("完成" if ok else "需要处理"))
        parts.append(f"{JOB_INFO[job]['label']}：{status}")
    return "；".join(parts) if parts else "暂无最近运行"


def _ai_template_draft_row(draft: dict[str, Any], *, product_name: str) -> dict[str, Any]:
    evidence = draft.get("evidence") if isinstance(draft.get("evidence"), dict) else {}
    differences = draft.get("differences_from_manual_template")
    risk_notes = draft.get("risk_notes")
    return {
        "草稿名": _text(draft.get("draft_name")),
        "产品": _text(product_name),
        "草稿 Key": _text(draft.get("draft_key")),
        "基于人工模板": _text(draft.get("base_manual_mode_key")),
        "候选素材": int(evidence.get("candidate_count") or 0),
        "消耗": round(float(evidence.get("stat_cost") or 0), 4),
        "转化": round(float(evidence.get("convert_cnt") or 0), 4),
        "ROI": round(float(evidence.get("avg_roi_1day") or 0), 4),
        "差异": "；".join(str(item) for item in differences if str(item).strip()) if isinstance(differences, list) else "",
        "风险": "；".join(str(item) for item in risk_notes if str(item).strip()) if isinstance(risk_notes, list) else "",
        "状态": _text(draft.get("status")),
    }


def _allowed_account_count(project_root: str | Path, path_text: str) -> tuple[int, str]:
    if not path_text:
        return 0, "未配置允许创建账户名单文件路径"
    path = Path(path_text)
    if not path.is_absolute():
        path = Path(project_root) / path
    if not path.exists():
        return 0, f"允许创建账户名单不存在：{path_text}"
    try:
        payload = load_json(path)
    except (FileNotFoundError, json.JSONDecodeError, ValueError):
        return 0, f"允许创建账户名单无法读取：{path_text}"
    rows = payload.get("allowed_target_accounts") if isinstance(payload.get("allowed_target_accounts"), list) else []
    count = 0
    for row in rows:
        if isinstance(row, dict) and _text(row.get("advertiser_id")) and _enabled(row.get("enable", row.get("enabled", True))):
            count += 1
    return count, ""


def _enabled_jobs(product: dict[str, Any]) -> list[str]:
    automation = product.get("automation") if isinstance(product.get("automation"), dict) else {}
    return [
        job
        for job in JOB_INFO
        if _enabled((automation.get(job) if isinstance(automation.get(job), dict) else {}).get("enabled"), False)
    ]


def _job_catalog() -> list[dict[str, str]]:
    return [{"value": job, "label": info["label"], "description": info["description"]} for job, info in JOB_INFO.items()]


def _source_account_name(product: dict[str, Any], account_names: dict[str, str]) -> str:
    explicit = _text(product.get("source_advertiser_name") or product.get("source_account_name"))
    if explicit:
        return explicit
    return account_name_for(account_names, product.get("source_advertiser_id"), fallback="未配置源素材账户名")


def _preload_target_scope(product: dict[str, Any]) -> str:
    automation = product.get("automation") if isinstance(product.get("automation"), dict) else {}
    preload = automation.get("source_material_preload") if isinstance(automation.get("source_material_preload"), dict) else {}
    return _text(preload.get("target_scope") or "allowed_accounts")


def _preload_scope_label(scope: str) -> str:
    if scope == "allowed_accounts":
        return "允许创建账户名单"
    if scope == "yesterday_spent":
        return "昨日有消耗账户"
    if scope == "today_spent":
        return "今日有消耗账户"
    return scope or "未配置"


def _account_keyword(product: dict[str, Any]) -> str:
    discovery = _account_discovery(product)
    return _text(discovery.get("account_name_keyword") or product.get("account_name_keyword") or product.get("product"))


def _account_remark(product: dict[str, Any]) -> str:
    discovery = _account_discovery(product)
    return _text(discovery.get("account_remark_equals") or product.get("account_remark_pattern"))


def _account_discovery(product: dict[str, Any]) -> dict[str, Any]:
    automation = product.get("automation") if isinstance(product.get("automation"), dict) else {}
    value = automation.get("account_discovery")
    return dict(value) if isinstance(value, dict) else {}


def _parse_allowed_accounts_upload(filename: str, content: bytes) -> list[dict[str, str]]:
    lower = filename.lower()
    if lower.endswith(".csv") or lower.endswith(".txt"):
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        return [_clean_upload_row(row) for row in reader if _has_upload_value(row)]
    if lower.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("当前环境缺少 openpyxl，无法读取 Excel 文件") from exc
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value or "").strip() for value in rows[0]]
        output = []
        for row in rows[1:]:
            record = {
                headers[index]: "" if value is None else str(value).strip()
                for index, value in enumerate(row)
                if index < len(headers) and headers[index]
            }
            if _has_upload_value(record):
                output.append(_clean_upload_row(record))
        return output
    raise RuntimeError("只支持 CSV、TXT 或 XLSX 文件")


def _clean_upload_row(row: dict[str, Any]) -> dict[str, str]:
    return {str(key or "").strip(): _text(value) for key, value in row.items()}


def _has_upload_value(row: dict[str, Any]) -> bool:
    return any(_text(value) for value in row.values())


def _field(row: dict[str, str], field: str) -> str:
    for alias in ALLOWED_ACCOUNT_HEADER_ALIASES[field]:
        if alias in row and _text(row.get(alias)):
            return _text(row.get(alias))
    return ""


def _enabled_for_import(value: str) -> bool:
    text = value.strip().lower()
    if not text:
        return True
    return text not in {"0", "false", "no", "n", "off", "disabled", "disable", "paused", "否", "不", "停用", "暂停"}


def _allowed_accounts_result(
    *,
    title: str,
    status: str,
    rows: list[dict[str, str]],
    counts: dict[str, int],
    blocking_reasons: list[str],
    artifact_path: str,
    raw: dict[str, Any],
) -> dict[str, Any]:
    warnings = []
    if status == "committed":
        warnings.append("只写入本地允许创建账户 JSON，未执行真实同步、补材或预推送动作。")
        warnings.append("页面会回填名单文件路径，点击保存配置后该产品才会引用这份名单。")
    return {
        "summary": {
            "title": title,
            "status": status,
            "risk_level": "medium" if blocking_reasons else "low",
            "execution_enabled": False,
            "items": [
                {"label": "读取行数", "value": counts.get("read", 0)},
                {"label": "启用账户", "value": counts.get("enabled", 0)},
                {"label": "跳过账户", "value": counts.get("skipped", 0)},
                {"label": "错误", "value": counts.get("error", 0)},
            ],
            "warnings": warnings,
            "blocking_reasons": blocking_reasons,
        },
        "table": {
            "columns": ["处理方式", "产品", "产品 Key", "账户名", "账户 ID", "是否启用", "备注", "问题"],
            "rows": rows,
        },
        "artifact_path": artifact_path,
        "raw": raw,
    }


def _blocked_result(title: str, reasons: list[str], raw: dict[str, Any]) -> dict[str, Any]:
    return {
        "summary": {
            "title": title,
            "status": "blocked",
            "risk_level": "medium",
            "execution_enabled": False,
            "items": [{"label": "问题数", "value": len(reasons)}],
            "warnings": [],
            "blocking_reasons": reasons,
        },
        "table": {"columns": ["问题"], "rows": [{"问题": reason} for reason in reasons]},
        "artifact_path": "",
        "raw": raw,
    }


def _enabled(value: Any, default: bool = True) -> bool:
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"1", "true", "yes", "y", "on", "enabled"}


def _project_path(project_root: str | Path, path_text: str) -> Path:
    path = Path(path_text)
    return path if path.is_absolute() else Path(project_root) / path


def _string_list(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [_text(item) for item in value if _text(item)]


def _parse_script_json(stdout: str) -> dict[str, Any]:
    try:
        value = json.loads(stdout)
    except json.JSONDecodeError:
        return {}
    return value if isinstance(value, dict) else {}


def _text(value: Any) -> str:
    return str(value or "").strip()
