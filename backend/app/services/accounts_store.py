from __future__ import annotations

import csv
import hashlib
import io
import json
from pathlib import Path
from typing import Any

from backend.app.services.atomic_write import write_json_atomic

ACCOUNT_FIELDS = [
    "product_key",
    "product_name",
    "advertiser_id",
    "advertiser_name",
    "channel",
    "owner",
    "account_remark",
    "status",
    "notes",
]

REQUIRED_FIELDS = ["product_key", "product_name", "advertiser_id", "advertiser_name"]
VALID_STATUSES = {"active", "paused", "disabled"}
CHANNEL_ALIASES = {"WECHAT_GAME": "微信", "wx": "微信", "wechat": "微信", "BYTEDANCE_GAME": "字节小游戏"}
BULK_UPDATE_FIELDS = {"channel": "渠道", "owner": "负责人", "account_remark": "备注", "status": "状态"}

_HISTORY_ACCOUNT_ID_KEYS = ("advertiser_id", "adv_id", "account_id", "operator", "target_advertiser_id")
_HISTORY_ADVERTISER_NAME_KEYS = ("advertiser_name", "account_name", "accountName")
_HISTORY_PRODUCT_KEY_KEYS = ("product_key", "productKey")
_HISTORY_PRODUCT_NAME_KEYS = ("product_name", "productName", "product", "brandName", "brand_name")
_HISTORY_CHANNEL_KEYS = ("channel", "platform")
_HISTORY_OWNER_KEYS = ("owner", "operator_name", "owner_name")
_HISTORY_REMARK_KEYS = ("account_remark", "remark")


def account_store_path(configs_dir: str | Path) -> Path:
    return Path(configs_dir) / "accounts" / "product-accounts.local.json"


def load_accounts(configs_dir: str | Path) -> list[dict[str, str]]:
    path = account_store_path(configs_dir)
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return []
    rows = value.get("accounts") if isinstance(value, dict) else []
    if not isinstance(rows, list):
        return []
    return [_normalize_account(row) for row in rows if isinstance(row, dict)]


def save_accounts(configs_dir: str | Path, accounts: list[dict[str, str]]) -> Path:
    path = account_store_path(configs_dir)
    payload = {"accounts": sorted(accounts, key=lambda row: (row["product_key"], row["advertiser_id"]))}
    return write_json_atomic(path, payload)


def parse_paste_text(text: str) -> list[dict[str, str]]:
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        return []
    delimiter = "\t" if "\t" in lines[0] else ","
    reader = csv.DictReader(lines, delimiter=delimiter)
    return [_normalize_account(row) for row in reader]


def parse_csv_bytes(content: bytes) -> list[dict[str, str]]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [_normalize_account(row) for row in reader]


def parse_xlsx_bytes(content: bytes) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("当前环境缺少 openpyxl，无法读取 Excel 文件") from exc
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value or "").strip() for value in rows[0]]
    output = []
    for row in rows[1:]:
        record = {headers[index]: "" if value is None else str(value).strip() for index, value in enumerate(row)}
        output.append(_normalize_account(record))
    return output


def filter_accounts(
    accounts: list[dict[str, str]],
    *,
    product_key: str = "",
    channel: str = "",
    owner: str = "",
    status: str = "",
    advertiser_ids: list[str] | None = None,
) -> list[dict[str, str]]:
    rows = accounts
    selected_ids = {str(item or "").strip() for item in (advertiser_ids or []) if str(item or "").strip()}
    if selected_ids:
        rows = [row for row in rows if row.get("advertiser_id") in selected_ids]
    if product_key:
        rows = [row for row in rows if row.get("product_key") == product_key]
    if channel:
        rows = [row for row in rows if row.get("channel") == channel]
    if owner:
        rows = [row for row in rows if row.get("owner") == owner]
    if status:
        rows = [row for row in rows if row.get("status") == status]
    return rows


def preview_import(existing: list[dict[str, str]], incoming: list[dict[str, str]]) -> dict[str, Any]:
    existing_by_id = {row["advertiser_id"]: row for row in existing}
    seen: set[str] = set()
    rows = []
    counts = {"read": len(incoming), "new": 0, "update": 0, "skip": 0, "error": 0}
    blocking_reasons = []
    for index, account in enumerate(incoming, start=1):
        errors = _validate_account(account)
        advertiser_id = account.get("advertiser_id", "")
        if advertiser_id in seen:
            errors.append("同次导入中账户 ID 重复")
        seen.add(advertiser_id)
        if errors:
            action = "错误"
            counts["error"] += 1
            blocking_reasons.extend(f"第 {index} 行：{error}" for error in errors)
        elif advertiser_id in existing_by_id:
            action = "更新"
            counts["update"] += 1
        else:
            action = "新增"
            counts["new"] += 1
        rows.append(_account_table_row(account, action, "；".join(errors)))

    status = "blocked" if counts["error"] else "planned"
    return {
        "summary": {
            "title": "账户导入预览",
            "status": status,
            "risk_level": "high" if counts["error"] else "low",
            "execution_enabled": not counts["error"],
            "items": [
                {"label": "读取行数", "value": counts["read"]},
                {"label": "新增", "value": counts["new"]},
                {"label": "更新", "value": counts["update"]},
                {"label": "跳过", "value": counts["skip"]},
                {"label": "错误", "value": counts["error"]},
            ],
            "warnings": [],
            "blocking_reasons": blocking_reasons,
        },
        "table": {"columns": _table_columns(), "rows": rows},
        "artifact_path": "",
        "raw": {"incoming": incoming, "counts": counts},
    }


def commit_import(configs_dir: str | Path, incoming: list[dict[str, str]]) -> dict[str, Any]:
    existing = load_accounts(configs_dir)
    preview = preview_import(existing, incoming)
    if preview["summary"]["blocking_reasons"]:
        return preview
    by_id = {row["advertiser_id"]: row for row in existing}
    for account in incoming:
        by_id[account["advertiser_id"]] = account
    path = save_accounts(configs_dir, list(by_id.values()))
    preview["summary"]["title"] = "账户导入完成"
    preview["summary"]["status"] = "committed"
    preview["summary"]["execution_enabled"] = False
    preview["artifact_path"] = str(path)
    return preview


def accounts_result(accounts: list[dict[str, str]], *, artifact_path: str) -> dict[str, Any]:
    rows = [_account_table_row(row, "", "") for row in accounts]
    for row in rows:
        row.pop("处理方式", None)
        row.pop("问题", None)
    return {
        "summary": {
            "title": "产品账户库",
            "status": "loaded",
            "risk_level": "low",
            "execution_enabled": False,
            "items": [{"label": "账户数", "value": len(rows)}],
            "warnings": [],
            "blocking_reasons": [],
        },
        "table": {
            "columns": ["产品", "产品 Key", "账户 ID", "账户名", "渠道", "负责人", "状态", "备注", "说明"],
            "rows": rows,
        },
        "artifact_path": artifact_path,
        "raw": {"accounts": accounts},
    }


def bulk_update_accounts(
    configs_dir: str | Path,
    *,
    filters: dict[str, str],
    updates: dict[str, str],
    advertiser_ids: list[str] | None = None,
) -> dict[str, Any]:
    preview = preview_bulk_update_accounts(configs_dir, filters=filters, updates=updates, advertiser_ids=advertiser_ids)
    if preview["summary"]["status"] == "blocked":
        return preview
    updated_accounts = preview.get("raw", {}).get("updated_accounts")
    if not isinstance(updated_accounts, list):
        return _bulk_update_blocked(filters, ["批量修改预览结果缺少待写入账户，已停止写入"])
    path = save_accounts(configs_dir, [dict(account) for account in updated_accounts if isinstance(account, dict)])
    preview["summary"]["title"] = "账户批量修改完成"
    preview["summary"]["status"] = "committed"
    preview["summary"]["execution_enabled"] = False
    preview["summary"]["warnings"] = ["已写入产品账户库；未命中账户保持不变。"]
    preview["artifact_path"] = str(path)
    return preview


def preview_bulk_update_accounts(
    configs_dir: str | Path,
    *,
    filters: dict[str, str],
    updates: dict[str, str],
    advertiser_ids: list[str] | None = None,
) -> dict[str, Any]:
    unsupported_fields = sorted(set(updates) - set(BULK_UPDATE_FIELDS))
    if unsupported_fields:
        return _bulk_update_blocked(
            filters,
            [f"不支持批量修改字段：{'、'.join(unsupported_fields)}"],
        )

    normalized_updates = {}
    for field, value in updates.items():
        text = str(value or "").strip()
        if field == "channel":
            normalized_updates[field] = _normalize_channel(text)
        elif field == "status":
            normalized_updates[field] = text
        else:
            normalized_updates[field] = text

    if not normalized_updates:
        return _bulk_update_blocked(filters, ["至少勾选一个要修改的字段"])
    if "status" in normalized_updates and normalized_updates["status"] not in VALID_STATUSES:
        return _bulk_update_blocked(filters, ["状态必须是 active、paused 或 disabled"])

    accounts = load_accounts(configs_dir)
    matched = filter_accounts(
        accounts,
        product_key=str(filters.get("product_key") or "").strip(),
        channel=_normalize_channel(str(filters.get("channel") or "").strip()),
        owner=str(filters.get("owner") or "").strip(),
        status=str(filters.get("status") or "").strip(),
        advertiser_ids=advertiser_ids,
    )
    if not matched:
        return _bulk_update_blocked(filters, ["当前筛选条件没有命中任何账户"])

    matched_ids = {row["advertiser_id"] for row in matched}
    updated_accounts = []
    updated_rows = []
    for account in accounts:
        if account["advertiser_id"] in matched_ids:
            next_account = dict(account)
            next_account.update(normalized_updates)
            normalized = _normalize_account(next_account)
            updated_accounts.append(normalized)
            updated_rows.append(_account_table_row(normalized, "已修改", ""))
        else:
            updated_accounts.append(account)

    changed_labels = "、".join(BULK_UPDATE_FIELDS[field] for field in normalized_updates)
    return {
        "summary": {
            "title": "账户批量修改预览",
            "status": "planned",
            "risk_level": "medium",
            "execution_enabled": True,
            "items": [
                {"label": "匹配账户", "value": len(updated_rows)},
                {"label": "修改字段", "value": changed_labels},
                {"label": "产品 Key 筛选", "value": str(filters.get("product_key") or "全部")},
                {"label": "渠道筛选", "value": str(filters.get("channel") or "全部")},
                {"label": "负责人筛选", "value": str(filters.get("owner") or "全部")},
                {"label": "状态筛选", "value": str(filters.get("status") or "全部")},
                {"label": "选中账户", "value": len({str(item or '').strip() for item in (advertiser_ids or []) if str(item or '').strip()}) or "未指定"},
            ],
            "warnings": ["这里只预览批量修改结果；点击确认写入后才会修改产品账户库。"],
            "blocking_reasons": [],
        },
        "table": {"columns": _table_columns(), "rows": updated_rows},
        "artifact_path": str(account_store_path(configs_dir)),
        "raw": {
            "filters": filters,
            "updates": normalized_updates,
            "matched_count": len(updated_rows),
            "updated_accounts": updated_accounts,
        },
    }


def accounts_to_csv(accounts: list[dict[str, str]]) -> str:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=ACCOUNT_FIELDS)
    writer.writeheader()
    for account in accounts:
        writer.writerow({field: account.get(field, "") for field in ACCOUNT_FIELDS})
    return buffer.getvalue()


def accounts_template_csv() -> str:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=ACCOUNT_FIELDS)
    writer.writeheader()
    writer.writerow(
        {
            "product_key": "diandian-hero",
            "product_name": "点点英雄",
            "advertiser_id": "1866125088740552",
            "advertiser_name": "黑旗游戏",
            "channel": "微信",
            "owner": "运营A",
            "account_remark": "点点英雄-黑旗",
            "status": "active",
            "notes": "示例：按这一行的表头填写，status 可填 active、paused、disabled",
        }
    )
    return buffer.getvalue()


def backfill_accounts_from_history(configs_dir: str | Path, runs_dir: str | Path) -> dict[str, Any]:
    existing = load_accounts(configs_dir)
    store_path = account_store_path(configs_dir)
    run_discovered, run_scanned_files = _discover_accounts_from_history(runs_dir)
    config_discovered, config_scanned_files = _discover_accounts_from_history(configs_dir, skip_paths={store_path})
    discovered = [*run_discovered, *config_discovered]
    scanned_files = run_scanned_files + config_scanned_files
    discovered_by_id = _dedupe_accounts(discovered)
    _apply_product_key_preferences(discovered_by_id, existing)
    _apply_existing_product_key_preferences(existing, discovered_by_id.values())
    existing_by_id = {row["advertiser_id"]: row for row in existing}
    merged_by_id = dict(existing_by_id)
    rows: list[dict[str, str]] = []
    counts = {
        "scanned_files": scanned_files,
        "discovered": len(discovered_by_id),
        "new": 0,
        "filled": 0,
        "unchanged": 0,
    }

    for advertiser_id, discovered_account in sorted(discovered_by_id.items()):
        existing_account = existing_by_id.get(advertiser_id)
        if existing_account is None:
            merged_by_id[advertiser_id] = discovered_account
            counts["new"] += 1
            rows.append(_account_table_row(discovered_account, "新增", ""))
            continue

        merged, changed_fields = _fill_missing_account_fields(existing_account, discovered_account)
        merged_by_id[advertiser_id] = merged
        if changed_fields:
            counts["filled"] += 1
            rows.append(_account_table_row(merged, "补全", "补全字段：" + "、".join(changed_fields)))
        else:
            counts["unchanged"] += 1
            rows.append(_account_table_row(existing_account, "保持", ""))

    path = save_accounts(configs_dir, list(merged_by_id.values())) if discovered_by_id else account_store_path(configs_dir)
    warnings = []
    if not discovered_by_id:
        warnings.append("没有从历史运行数据中识别到账户和产品信息。")
    return {
        "summary": {
            "title": "历史数据补全账户库",
            "status": "committed" if discovered_by_id else "empty",
            "risk_level": "low",
            "execution_enabled": False,
            "items": [
                {"label": "扫描 JSON 文件", "value": counts["scanned_files"]},
                {"label": "历史识别账户", "value": counts["discovered"]},
                {"label": "新增", "value": counts["new"]},
                {"label": "补全", "value": counts["filled"]},
                {"label": "保持不变", "value": counts["unchanged"]},
            ],
            "warnings": warnings,
            "blocking_reasons": [],
        },
        "table": {"columns": _table_columns(), "rows": rows},
        "artifact_path": str(path),
        "raw": {"counts": counts, "accounts": list(discovered_by_id.values())},
    }


def _normalize_account(row: dict[str, Any]) -> dict[str, str]:
    normalized = {field: str(row.get(field) or "").strip() for field in ACCOUNT_FIELDS}
    normalized["channel"] = _normalize_channel(normalized["channel"])
    normalized["status"] = normalized["status"] or "active"
    return normalized


def _validate_account(account: dict[str, str]) -> list[str]:
    errors = [f"缺少 {field}" for field in REQUIRED_FIELDS if not account.get(field)]
    if account.get("status") not in VALID_STATUSES:
        errors.append("status 必须是 active、paused 或 disabled")
    return errors


def _account_table_row(account: dict[str, str], action: str, issue: str) -> dict[str, str]:
    return {
        "处理方式": action,
        "产品": account.get("product_name", ""),
        "产品 Key": account.get("product_key", ""),
        "账户 ID": account.get("advertiser_id", ""),
        "账户名": account.get("advertiser_name", ""),
        "渠道": account.get("channel", ""),
        "负责人": account.get("owner", ""),
        "状态": account.get("status", ""),
        "备注": account.get("account_remark", ""),
        "说明": account.get("notes", ""),
        "问题": issue,
    }


def _table_columns() -> list[str]:
    return ["处理方式", "产品", "产品 Key", "账户 ID", "账户名", "渠道", "负责人", "状态", "备注", "说明", "问题"]


def _bulk_update_blocked(filters: dict[str, str], reasons: list[str]) -> dict[str, Any]:
    return {
        "summary": {
            "title": "账户批量修改被拦截",
            "status": "blocked",
            "risk_level": "low",
            "execution_enabled": False,
            "items": [
                {"label": "产品 Key 筛选", "value": str(filters.get("product_key") or "全部")},
                {"label": "渠道筛选", "value": str(filters.get("channel") or "全部")},
                {"label": "负责人筛选", "value": str(filters.get("owner") or "全部")},
                {"label": "状态筛选", "value": str(filters.get("status") or "全部")},
            ],
            "warnings": [],
            "blocking_reasons": reasons,
        },
        "table": {"columns": _table_columns(), "rows": []},
        "artifact_path": "",
        "raw": {"filters": filters, "blocking_reasons": reasons},
    }


def _discover_accounts_from_history(
    runs_dir: str | Path,
    *,
    skip_paths: set[Path] | None = None,
) -> tuple[list[dict[str, str]], int]:
    root = Path(runs_dir)
    if not root.exists():
        return [], 0
    skipped = {path.resolve() for path in (skip_paths or set())}
    accounts: list[dict[str, str]] = []
    scanned_files = 0
    for path in sorted(root.rglob("*.json")):
        if path.resolve() in skipped:
            continue
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError, UnicodeDecodeError):
            continue
        scanned_files += 1
        _collect_accounts_from_value(payload, {}, accounts)
    return accounts, scanned_files


def _collect_accounts_from_value(value: Any, context: dict[str, str], output: list[dict[str, str]]) -> None:
    if isinstance(value, list):
        for item in value:
            _collect_accounts_from_value(item, context, output)
        return
    if not isinstance(value, dict):
        return

    next_context = dict(context)
    next_context.update({key: value for key, value in _context_from_node(value).items() if value})
    for embedded_key in ("create_request", "request", "source_request"):
        embedded = value.get(embedded_key)
        if isinstance(embedded, dict):
            next_context.update({key: value for key, value in _context_from_node(embedded).items() if value})

    account = _account_from_node(value, next_context)
    if account:
        output.append(account)

    for child in value.values():
        _collect_accounts_from_value(child, next_context, output)


def _context_from_node(node: dict[str, Any]) -> dict[str, str]:
    return {
        "product_key": _pick_text(node, _HISTORY_PRODUCT_KEY_KEYS),
        "product_name": _pick_text(node, _HISTORY_PRODUCT_NAME_KEYS),
        "channel": _pick_text(node, _HISTORY_CHANNEL_KEYS),
        "owner": _pick_text(node, _HISTORY_OWNER_KEYS),
    }


def _account_from_node(node: dict[str, Any], context: dict[str, str]) -> dict[str, str] | None:
    advertiser_id = _pick_text(node, _HISTORY_ACCOUNT_ID_KEYS)
    if not advertiser_id or not advertiser_id.isdigit():
        return None

    product_name = _pick_text(node, _HISTORY_PRODUCT_NAME_KEYS) or context.get("product_name", "")
    product_key = _pick_text(node, _HISTORY_PRODUCT_KEY_KEYS) or context.get("product_key", "")
    if not product_key and product_name:
        product_key = _derive_product_key(node, product_name)
    if not product_name and product_key:
        product_name = product_key
    if not product_key or not product_name:
        return None

    account = {
        "product_key": product_key,
        "product_name": product_name,
        "advertiser_id": advertiser_id,
        "advertiser_name": _pick_text(node, _HISTORY_ADVERTISER_NAME_KEYS) or advertiser_id,
        "channel": _pick_text(node, _HISTORY_CHANNEL_KEYS) or context.get("channel", "") or _infer_channel(node),
        "owner": _pick_text(node, _HISTORY_OWNER_KEYS) or context.get("owner", ""),
        "account_remark": _pick_text(node, _HISTORY_REMARK_KEYS),
        "status": "active",
        "notes": "历史数据自动补全",
    }
    return _normalize_account(account)


def _dedupe_accounts(accounts: list[dict[str, str]]) -> dict[str, dict[str, str]]:
    by_id: dict[str, dict[str, str]] = {}
    for account in accounts:
        advertiser_id = account.get("advertiser_id", "")
        if not advertiser_id:
            continue
        existing = by_id.get(advertiser_id)
        by_id[advertiser_id] = account if existing is None else _merge_discovered_account(existing, account)
    return by_id


def _apply_product_key_preferences(discovered_by_id: dict[str, dict[str, str]], existing: list[dict[str, str]]) -> None:
    preferred_by_product: dict[str, str] = {}
    for account in [*existing, *discovered_by_id.values()]:
        product_name = account.get("product_name", "")
        product_key = account.get("product_key", "")
        if product_name and product_key and not product_key.startswith("history-"):
            preferred_by_product.setdefault(product_name, product_key)

    for account in discovered_by_id.values():
        preferred = preferred_by_product.get(account.get("product_name", ""))
        if preferred and account.get("product_key", "").startswith("history-"):
            account["product_key"] = preferred


def _apply_existing_product_key_preferences(
    existing: list[dict[str, str]], discovered: Any
) -> None:
    preferred_by_product: dict[str, str] = {}
    for account in [*existing, *list(discovered)]:
        product_name = account.get("product_name", "")
        product_key = account.get("product_key", "")
        if product_name and product_key and not product_key.startswith("history-"):
            preferred_by_product.setdefault(product_name, product_key)

    for account in existing:
        preferred = preferred_by_product.get(account.get("product_name", ""))
        if preferred and account.get("product_key", "").startswith("history-"):
            account["product_key"] = preferred


def _merge_discovered_account(left: dict[str, str], right: dict[str, str]) -> dict[str, str]:
    merged = dict(left)
    for field in ACCOUNT_FIELDS:
        if _should_replace_account_field(merged, right, field):
            merged[field] = right[field]
    return _normalize_account(merged)


def _fill_missing_account_fields(existing: dict[str, str], discovered: dict[str, str]) -> tuple[dict[str, str], list[str]]:
    merged = dict(existing)
    changed_fields = []
    for field in ACCOUNT_FIELDS:
        if field == "status":
            continue
        if _should_replace_account_field(merged, discovered, field):
            merged[field] = discovered[field]
            changed_fields.append(field)
    return _normalize_account(merged), changed_fields


def _should_replace_account_field(left: dict[str, str], right: dict[str, str], field: str) -> bool:
    if not right.get(field):
        return False
    if not left.get(field):
        return True
    if field == "advertiser_name":
        left_name = left.get("advertiser_name", "")
        right_name = right.get("advertiser_name", "")
        advertiser_id = left.get("advertiser_id") or right.get("advertiser_id", "")
        return left_name == advertiser_id and right_name != advertiser_id
    return False


def _pick_text(node: dict[str, Any], keys: tuple[str, ...]) -> str:
    for key in keys:
        value = node.get(key)
        if value is None:
            continue
        text = str(value).strip()
        if text:
            return text
    return ""


def _derive_product_key(node: dict[str, Any], product_name: str) -> str:
    brand_id = str(node.get("brandId") or node.get("brand_id") or "").strip()
    if brand_id:
        return f"brand-{brand_id}"
    digest = hashlib.sha1(product_name.encode("utf-8")).hexdigest()[:8]
    return f"history-{digest}"


def _infer_channel(node: dict[str, Any]) -> str:
    username = str(node.get("username") or "").strip()
    if username.startswith("gh_"):
        return "微信"
    store_map = node.get("external_app_store_developers_map")
    if isinstance(store_map, dict):
        if "微信" in store_map:
            return "微信"
        if "抖音" in store_map:
            return "抖音"
    return ""


def _normalize_channel(value: str) -> str:
    return CHANNEL_ALIASES.get(value, value)
