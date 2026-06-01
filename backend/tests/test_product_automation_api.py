import csv
import io
import json
import sqlite3
from pathlib import Path

from fastapi.testclient import TestClient

from backend.app.main import create_app
from roibang_v2.db.bootstrap import bootstrap_database


def _client(tmp_path: Path) -> TestClient:
    return TestClient(create_app(project_root=tmp_path))


def _write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def _write_product_fixture(tmp_path: Path) -> None:
    _write_json(
        tmp_path / "configs" / "allowed-create-accounts.demo.local.json",
        {
            "allowed_target_accounts": [
                {"advertiser_id": "1001", "account_name": "演示账户一", "enable": True},
                {"advertiser_id": "1002", "account_name": "演示账户二", "enable": False},
            ]
        },
    )
    _write_json(
        tmp_path / "configs" / "products" / "demo-game.local.json",
        {
            "product_key": "demo-game",
            "product": "演示游戏",
            "platform": "WECHAT_GAME",
            "source_advertiser_id": "source-1",
            "source_advertiser_name": "演示源素材账户",
            "organization_id": "org-1",
            "allowed_target_accounts_path": "configs/allowed-create-accounts.demo.local.json",
            "automation": {
                "enabled": True,
                "account_discovery": {
                    "account_name_keyword": "演示游戏",
                    "account_remark_equals": "演示游戏-微小",
                },
                "source_material_preload": {"enabled": True, "target_scope": "allowed_accounts"},
                "source_material_rollup": {"enabled": True},
            },
        },
    )


def _write_create_mode_fixture(tmp_path: Path) -> None:
    mode_dir = tmp_path / "configs" / "create-modes"
    mode_dir.mkdir(parents=True, exist_ok=True)
    _write_json(
        mode_dir / "wx_pay_general_recent_scale.example.json",
        {
            "mode_key": "wx_pay_general_recent_scale",
            "display_name": "每付通投近期放量",
            "template_key": "wx_pay_general",
            "template_name_suffix": "每付通投近期放量",
            "defaults": {
                "daily_budget": 10000,
                "cpa_bid": 111,
                "project_count": 5,
                "units_per_project": 1,
            },
            "material_requirements": {
                "materials_per_unit": 6,
                "allow_reuse_on_insufficient": True,
            },
            "material_selection": {
                "lookback_days": 7,
                "selection_type": "high_spend",
                "random_shuffle": True,
            },
        },
    )


def _write_template_draft_db_fixture(tmp_path: Path) -> None:
    db_path = tmp_path / "data" / "roibang_v2.sqlite3"
    bootstrap_database(db_path)
    with sqlite3.connect(db_path) as conn:
        for index in range(35):
            conn.execute(
                """
                INSERT INTO product_source_material_metric_rollups (
                  product, source_advertiser_id, organization_id, window_key, window_days,
                  period_start, period_end, material_id, material_type, source_video_id,
                  name, review_status, signature, duration, file_size, create_time,
                  first_seen_metric_date, effective_create_date, effective_create_date_source,
                  tag_ids_json, account_count, project_count, promotion_count, stat_cost,
                  show_cnt, click_cnt, convert_cnt, active_register, roi_1day_cost_weighted,
                  roi_7days_cost_weighted, source, synced_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    "演示游戏",
                    "source-1",
                    "org-1",
                    "last_7d",
                    7,
                    "2026-05-22",
                    "2026-05-28",
                    f"material-{index}",
                    "video",
                    f"video-{index:03d}",
                    f"演示素材{index}",
                    "审核通过",
                    "",
                    0,
                    0,
                    "2026-05-28",
                    "2026-05-28",
                    "2026-05-28",
                    "first_seen_metric_date",
                    "[]",
                    1,
                    1,
                    1,
                    600,
                    6000,
                    600,
                    2,
                    2,
                    0.06,
                    0.08,
                    "unit_test",
                    "now",
                ),
            )


def test_product_automation_overview_lists_product_configs_with_chinese_labels(tmp_path: Path):
    _write_product_fixture(tmp_path)
    client = _client(tmp_path)

    response = client.get("/api/product-automation/overview")

    assert response.status_code == 200
    payload = response.json()
    assert payload["summary"]["title"] == "产品自动化配置"
    assert payload["summary"]["execution_enabled"] is False
    row = payload["table"]["rows"][0]
    assert row["产品"] == "演示游戏"
    assert row["源素材账户名"] == "演示源素材账户"
    assert row["源素材账户 ID"] == "source-1"
    assert row["允许创建账户"] == 1
    assert row["预推送目标"] == "允许创建账户名单"
    assert "源素材预推送" in row["启用定时任务"]
    assert payload["raw"]["jobs"][0]["label"] == "每日素材明细同步"


def test_product_automation_save_defaults_preload_to_allowed_accounts(tmp_path: Path):
    client = _client(tmp_path)
    body = {
        "product_key": "new-game",
        "product": "新游戏",
        "platform": "WECHAT_GAME",
        "source_advertiser_name": "新游戏源素材账户",
        "source_advertiser_id": "source-2",
        "organization_id": "org-2",
        "allowed_target_accounts_path": "configs/allowed-create-accounts.new.local.json",
        "account_name_keyword": "新游戏",
        "account_remark_equals": "新游戏-微小",
        "enabled_jobs": ["material_daily_sync", "source_material_preload"],
    }

    response = client.post("/api/product-automation/config/save", json=body)

    assert response.status_code == 200
    payload = response.json()
    assert payload["summary"]["status"] == "committed"
    saved = json.loads((tmp_path / "configs/products/new-game.local.json").read_text(encoding="utf-8"))
    assert saved["source_advertiser_name"] == "新游戏源素材账户"
    assert saved["automation"]["source_material_preload"]["enabled"] is True
    assert saved["automation"]["source_material_preload"]["target_scope"] == "allowed_accounts"
    assert saved["automation"]["material_daily_sync"]["enabled"] is True


def test_product_automation_allowed_accounts_template_downloads_xlsx(tmp_path: Path):
    from openpyxl import load_workbook

    client = _client(tmp_path)

    response = client.get(
        "/api/product-automation/allowed-accounts/template",
        params={"product_key": "demo-game", "product": "演示游戏"},
    )

    assert response.status_code == 200
    assert "allowed-create-accounts-template.demo-game.xlsx" in response.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(response.content), read_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    assert headers == ["产品名", "产品 Key", "账户名", "账户 ID", "是否启用", "备注"]
    first_row = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
    assert first_row[0] == "演示游戏"
    assert first_row[1] == "demo-game"


def test_product_automation_allowed_accounts_upload_writes_json_and_returns_path(tmp_path: Path):
    client = _client(tmp_path)
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=["产品名", "产品 Key", "账户名", "账户 ID", "是否启用", "备注"])
    writer.writeheader()
    writer.writerow({"产品名": "演示游戏", "产品 Key": "demo-game", "账户名": "演示账户一", "账户 ID": "1001", "是否启用": "是", "备注": "郭靖账户"})
    writer.writerow({"产品名": "演示游戏", "产品 Key": "demo-game", "账户名": "演示账户二", "账户 ID": "1002", "是否启用": "否", "备注": "暂停"})

    response = client.post(
        "/api/product-automation/allowed-accounts/import",
        params={"product_key": "demo-game", "product": "演示游戏", "platform": "WECHAT_GAME"},
        files={"file": ("allowed.csv", buffer.getvalue().encode("utf-8"), "text/csv")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["summary"]["title"] == "允许创建账户名单已生成"
    assert {"label": "启用账户", "value": 1} in payload["summary"]["items"]
    assert payload["raw"]["allowed_target_accounts_path"] == "configs/allowed-create-accounts.demo-game.local.json"
    saved_path = tmp_path / "configs" / "allowed-create-accounts.demo-game.local.json"
    saved = json.loads(saved_path.read_text(encoding="utf-8"))
    assert saved["product"] == "演示游戏"
    assert saved["allowed_target_accounts"] == [
        {
            "account_name": "演示账户一",
            "account_remark": "郭靖账户",
            "advertiser_id": "1001",
            "enable": True,
            "platform": "WECHAT_GAME",
            "product": "演示游戏",
            "product_key": "demo-game",
        }
    ]


def test_product_automation_dry_run_writes_request_without_execute(tmp_path: Path):
    _write_product_fixture(tmp_path)
    client = _client(tmp_path)

    response = client.post(
        "/api/product-automation/dry-run",
        json={"product_key": "demo-game", "job": "source_material_preload", "target_date": "yesterday"},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["summary"]["title"] == "产品自动化预演"
    assert payload["summary"]["execution_enabled"] is False
    assert payload["table"]["rows"][0]["真实执行"] == "否"
    command = payload["table"]["rows"][0]["脚本命令"]
    assert "--execute" not in command
    assert "--yes" not in command
    request_path = Path(payload["table"]["rows"][0]["请求 JSON"])
    assert request_path.exists()
    request_payload = json.loads(request_path.read_text(encoding="utf-8"))
    assert request_payload["source_material_preload_to_accounts"]["target_accounts"]["accounts"][0]["account_name"] == "演示账户一"


def test_product_automation_ai_template_drafts_returns_readonly_drafts(tmp_path: Path):
    _write_product_fixture(tmp_path)
    _write_create_mode_fixture(tmp_path)
    _write_template_draft_db_fixture(tmp_path)
    client = _client(tmp_path)

    response = client.post(
        "/api/product-automation/ai-template-drafts",
        json={"product_key": "demo-game", "max_drafts": 2},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["summary"]["title"] == "AI 模板草稿"
    assert payload["summary"]["status"] == "draft_only"
    assert payload["summary"]["execution_enabled"] is False
    assert {"label": "真实执行", "value": "否"} in payload["summary"]["items"]
    assert payload["table"]["columns"] == [
        "草稿名",
        "产品",
        "草稿 Key",
        "基于人工模板",
        "候选素材",
        "消耗",
        "转化",
        "ROI",
        "差异",
        "风险",
        "状态",
    ]
    row = payload["table"]["rows"][0]
    assert row["产品"] == "演示游戏"
    assert row["状态"] == "draft_only"
    assert row["基于人工模板"] == "wx_pay_general_recent_scale"
    assert row["候选素材"] == 35
    assert Path(payload["artifact_path"]).exists()
    assert payload["raw"]["manual_template_boundary"]["writes_manual_template"] is False
    assert payload["raw"]["manual_template_boundary"]["generates_create_plan"] is False
    assert payload["raw"]["actions"] == []


def test_product_automation_ai_template_draft_preview_writes_preview_only_json(tmp_path: Path):
    _write_product_fixture(tmp_path)
    _write_create_mode_fixture(tmp_path)
    _write_template_draft_db_fixture(tmp_path)
    client = _client(tmp_path)
    drafts_response = client.post(
        "/api/product-automation/ai-template-drafts",
        json={"product_key": "demo-game", "max_drafts": 2},
    )
    drafts_payload = drafts_response.json()
    draft_key = drafts_payload["table"]["rows"][0]["草稿 Key"]

    response = client.post(
        "/api/product-automation/ai-template-draft-preview",
        json={
            "product_key": "demo-game",
            "artifact_path": drafts_payload["artifact_path"],
            "draft_key": draft_key,
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["summary"]["title"] == "AI 模板草稿转正预览"
    assert payload["summary"]["status"] == "planned"
    assert payload["summary"]["execution_enabled"] is False
    assert {"label": "真实执行", "value": "否"} in payload["summary"]["items"]
    row = payload["table"]["rows"][0]
    assert row["产品"] == "演示游戏"
    assert row["草稿 Key"] == draft_key
    assert row["状态"] == "preview_only"
    target_path = tmp_path / "configs" / "create-modes" / "demo-game" / f"{draft_key}.local.json"
    assert row["目标文件"] == str(target_path)
    assert Path(payload["artifact_path"]).exists()
    assert not target_path.exists()
    assert payload["raw"]["preview"]["execution"]["enabled"] is False
    assert payload["raw"]["preview"]["proposed_create_mode"]["product"] == "演示游戏"


def test_product_automation_ai_template_draft_promote_calls_fixed_script_and_writes_mode(tmp_path: Path):
    _write_product_fixture(tmp_path)
    _write_create_mode_fixture(tmp_path)
    _write_template_draft_db_fixture(tmp_path)
    client = _client(tmp_path)
    drafts_response = client.post(
        "/api/product-automation/ai-template-drafts",
        json={"product_key": "demo-game", "max_drafts": 2},
    )
    drafts_payload = drafts_response.json()
    draft_key = drafts_payload["table"]["rows"][0]["草稿 Key"]
    preview_response = client.post(
        "/api/product-automation/ai-template-draft-preview",
        json={
            "product_key": "demo-game",
            "artifact_path": drafts_payload["artifact_path"],
            "draft_key": draft_key,
        },
    )
    preview_payload = preview_response.json()

    response = client.post(
        "/api/product-automation/ai-template-draft-promote",
        json={
            "product_key": "demo-game",
            "preview_path": preview_payload["artifact_path"],
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["summary"]["title"] == "AI 模板草稿已写入创建模式"
    assert payload["summary"]["status"] == "committed"
    assert payload["summary"]["execution_enabled"] is False
    assert {"label": "真实执行", "value": "否"} in payload["summary"]["items"]
    target_path = tmp_path / "configs" / "create-modes" / "demo-game" / f"{draft_key}.local.json"
    assert target_path.exists()
    saved = json.loads(target_path.read_text(encoding="utf-8"))
    assert saved["product"] == "演示游戏"
    assert saved["product_key"] == "demo-game"
    assert saved["mode_key"] == draft_key
    assert payload["table"]["rows"][0]["固定脚本"] == "scripts/run_ai_template_draft_promote.py"
    assert payload["table"]["rows"][0]["真实执行"] == "否"
    assert payload["raw"]["result"]["execution_enabled"] is False
    assert payload["raw"]["result"]["actions"] == []
